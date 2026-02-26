"""DOCX resume parser using python-docx.

Detects section headers via paragraph styles, bold runs,
font size, ALL-CAPS, and known-header pattern matching.
Also extracts content from tables (common in layout-based templates).
"""

from __future__ import annotations

from docx import Document

from src.constants import MAX_HEADER_CHARS, MAX_HEADER_WORDS
from src.parsers.base import BaseResumeParser


class DOCXResumeParser(BaseResumeParser):
    """Parse a **DOCX** resume into ``[{section, content}, …]``."""

    def parse(self) -> list[dict[str, str]]:
        doc = Document(str(self.file_path))

        entries: list[tuple[str, bool]] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            entries.append((text, self._is_heading(para, text)))

        for table in doc.tables:
            for row in table.rows:
                seen: set[str] = set()
                for cell in row.cells:
                    ct = cell.text.strip()
                    if ct and ct not in seen:
                        seen.add(ct)
                        is_hdr = (
                            self.matches_known_header(ct)
                            and len(ct.split()) <= MAX_HEADER_WORDS
                        )
                        entries.append((ct, is_hdr))

        return self._collect_sections(entries)

    # ── heading detection ────────────────────────────────────────

    @staticmethod
    def _is_heading(para, text: str) -> bool:
        if not text or len(text) > MAX_HEADER_CHARS:
            return False

        style_name = (para.style.name or "").lower()
        if "heading" in style_name or "title" in style_name:
            return True

        runs = [r for r in para.runs if r.text.strip()]
        all_bold = bool(runs) and all(r.bold for r in runs)
        has_large_font = any(
            r.font.size and r.font.size.pt >= 13
            for r in runs
            if r.font.size
        )
        is_upper = (
            text == text.upper()
            and len(text) > 2
            and any(c.isalpha() for c in text)
        )
        known = BaseResumeParser.matches_known_header(text)
        short = len(text.split()) <= MAX_HEADER_WORDS

        if known and (all_bold or has_large_font or is_upper):
            return True
        if all_bold and has_large_font and short:
            return True
        if is_upper and known:
            return True
        if all_bold and is_upper and short:
            return True
        return False
